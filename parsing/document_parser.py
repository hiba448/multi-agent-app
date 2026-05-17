import fitz  # PyMuPDF
from pptx import Presentation
from typing import List, Dict
import os


def parse_pdf(file_path: str) -> List[Dict]:
    """Extract text from a PDF file, page by page."""
    chunks = []
    doc = fitz.open(file_path)

    for page_num, page in enumerate(doc, start=1):
        text = page.get_text().strip()
        if text:
            chunks.append({
                "page": page_num,
                "content": text
            })

    doc.close()
    return chunks


def parse_pptx(file_path: str) -> List[Dict]:
    """Extract text from a PowerPoint file, slide by slide."""
    chunks = []
    prs = Presentation(file_path)

    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_text = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        slide_text.append(text)

        if slide_text:
            chunks.append({
                "page": slide_num,
                "content": "\n".join(slide_text)
            })

    return chunks


def split_into_chunks(
    pages: List[Dict],
    chunk_size: int = 500,
    overlap: int = 50
) -> List[Dict]:
    """
    Split page-level text into smaller overlapping chunks.
    chunk_size: approximate number of words per chunk.
    overlap: number of words shared between consecutive chunks.
    """
    all_chunks = []

    for page in pages:
        words = page["content"].split()
        start = 0

        while start < len(words):
            end = start + chunk_size
            chunk_words = words[start:end]
            chunk_text = " ".join(chunk_words)

            all_chunks.append({
                "page": page["page"],
                "content": chunk_text
            })

            if end >= len(words):
                break

            start += chunk_size - overlap

    return all_chunks


def parse_document(file_path: str) -> List[Dict]:
    """
    Main entry point. Detects file type, parses it,
    and returns a list of overlapping text chunks.
    """
    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":
        pages = parse_pdf(file_path)
    elif extension in [".pptx", ".ppt"]:
        pages = parse_pptx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {extension}")

    chunks = split_into_chunks(pages)
    return chunks